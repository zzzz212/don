#!/usr/bin/env python3
"""Презентация-отчёт по курсовой работе"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# МГИМО синий + классический академический стиль
MGIMO_BLUE = RGBColor(0x00, 0x2D, 0x6E)
MGIMO_LIGHT = RGBColor(0x1A, 0x4D, 0x9E)
ACCENT = RGBColor(0xC8, 0x9B, 0x3C)  # МГИМО золото
LIGHT_BG = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK = RGBColor(0x1A, 0x1F, 0x36)
TEXT_GREY = RGBColor(0x5C, 0x6B, 0x7C)
GREEN = RGBColor(0x10, 0x99, 0x6F)
BORDER = RGBColor(0xD5, 0xDB, 0xE4)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SW = prs.slide_width
SH = prs.slide_height
BLANK = prs.slide_layouts[6]


def add_bg(slide, color=WHITE):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg


def add_text(slide, left, top, width, height, text, *, font_size=18,
             bold=False, color=TEXT_DARK, align=PP_ALIGN.LEFT,
             font_name="Calibri", anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return tb


def add_rect(slide, left, top, width, height, fill, line_color=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_color:
        sh.line.color.rgb = line_color
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_round(slide, left, top, width, height, fill, line_color=None,
              corner=0.1):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width,
                                height)
    sh.adjustments[0] = corner
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_color:
        sh.line.color.rgb = line_color
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_oval(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_footer(slide, page_num, total):
    add_rect(slide, Inches(0.6), SH - Inches(0.45), Inches(12.1), Inches(0.02),
             BORDER)
    add_text(slide, Inches(0.6), SH - Inches(0.4), Inches(8), Inches(0.3),
             "Курсовая работа  |  МГИМО МИД России  |  Хадызов З.Д.",
             font_size=9, color=TEXT_GREY)
    add_text(slide, SW - Inches(2), SH - Inches(0.4), Inches(1.4), Inches(0.3),
             f"стр. {page_num} из {total}", font_size=9, color=TEXT_GREY,
             align=PP_ALIGN.RIGHT)


def header(slide, num, title, subtitle=None):
    # Номер раздела + название
    add_rect(slide, Inches(0.6), Inches(0.4), Inches(0.08), Inches(0.9),
             ACCENT)
    add_text(slide, Inches(0.85), Inches(0.4), Inches(1.2), Inches(0.4),
             num, font_size=12, color=ACCENT, bold=True)
    add_text(slide, Inches(0.85), Inches(0.75), Inches(11), Inches(0.7),
             title, font_size=28, bold=True, color=MGIMO_BLUE)
    if subtitle:
        add_text(slide, Inches(0.85), Inches(1.4), Inches(11), Inches(0.4),
                 subtitle, font_size=14, color=TEXT_GREY, italic=True)
    # Разделительная линия
    add_rect(slide, Inches(0.6), Inches(1.85), Inches(12.1), Inches(0.02),
             BORDER)


TOTAL = 16

# ==================== СЛАЙД 1: ТИТУЛЬНЫЙ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s, WHITE)

# Верхняя полоска
add_rect(s, 0, 0, SW, Inches(0.15), MGIMO_BLUE)
add_rect(s, 0, Inches(0.15), SW, Inches(0.04), ACCENT)

# Шапка с университетом
add_text(s, Inches(0.6), Inches(0.5), Inches(12.1), Inches(0.5),
         "МИНИСТЕРСТВО ИНОСТРАННЫХ ДЕЛ РОССИЙСКОЙ ФЕДЕРАЦИИ",
         font_size=12, color=MGIMO_BLUE, align=PP_ALIGN.CENTER, bold=True)
add_text(s, Inches(0.6), Inches(0.95), Inches(12.1), Inches(0.5),
         "Московский государственный институт международных отношений (МГИМО)",
         font_size=14, color=MGIMO_BLUE, align=PP_ALIGN.CENTER, bold=True)
add_text(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.4),
         "Факультет управления и политики",
         font_size=12, color=TEXT_DARK, align=PP_ALIGN.CENTER)

# Разделитель
add_rect(s, Inches(4), Inches(2.0), Inches(5.3), Inches(0.03), ACCENT)

# Тип работы
add_text(s, Inches(0.6), Inches(2.3), Inches(12.1), Inches(0.4),
         "КУРСОВАЯ РАБОТА",
         font_size=14, color=ACCENT, align=PP_ALIGN.CENTER, bold=True)

# Тема работы
add_text(s, Inches(0.6), Inches(3.0), Inches(12.1), Inches(0.4),
         "на тему:",
         font_size=14, color=TEXT_GREY, align=PP_ALIGN.CENTER, italic=True)

add_round(s, Inches(0.8), Inches(3.5), Inches(11.7), Inches(1.6), MGIMO_BLUE)
add_text(s, Inches(1.0), Inches(3.7), Inches(11.3), Inches(0.7),
         "Разработка цифровой платформы",
         font_size=24, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(4.2), Inches(11.3), Inches(0.6),
         "правовой поддержки малого и среднего бизнеса",
         font_size=20, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(4.65), Inches(11.3), Inches(0.4),
         "с использованием технологий искусственного интеллекта",
         font_size=16, color=WHITE, align=PP_ALIGN.CENTER, italic=True)

# Автор и научный руководитель
add_text(s, Inches(0.6), Inches(5.5), Inches(6.5), Inches(0.4),
         "Выполнил:", font_size=12, color=TEXT_GREY, align=PP_ALIGN.RIGHT)
add_text(s, Inches(7.2), Inches(5.5), Inches(5.5), Inches(0.4),
         "Хадызов Зураб Джамалаевич",
         font_size=14, color=TEXT_DARK, bold=True)
add_text(s, Inches(7.2), Inches(5.85), Inches(5.5), Inches(0.4),
         "студент 3 курса",
         font_size=12, color=TEXT_GREY)

add_text(s, Inches(0.6), Inches(6.3), Inches(6.5), Inches(0.4),
         "Научный руководитель:", font_size=12, color=TEXT_GREY,
         align=PP_ALIGN.RIGHT)
add_text(s, Inches(7.2), Inches(6.3), Inches(5.5), Inches(0.4),
         "________________________",
         font_size=12, color=TEXT_DARK)

# Город и год
add_text(s, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.3),
         "Москва · 2026",
         font_size=12, color=MGIMO_BLUE, align=PP_ALIGN.CENTER, bold=True)

# ==================== СЛАЙД 2: СОДЕРЖАНИЕ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "СОДЕРЖАНИЕ", "Структура работы")

contents = [
    ("01", "Введение и актуальность"),
    ("02", "Цель и задачи исследования"),
    ("03", "Объект, предмет и методология"),
    ("04", "Анализ предметной области"),
    ("05", "Архитектура системы"),
    ("06", "Технологический стек"),
    ("07", "Реализация: AI-анализ документов"),
    ("08", "Реализация: проверка контрагентов"),
    ("09", "Реализация: база права и генерация"),
    ("10", "Деплой и инфраструктура"),
    ("11", "Результаты и достижения"),
    ("12", "Сложности разработки"),
    ("13", "Перспективы развития"),
    ("14", "Заключение"),
]

x_left = Inches(0.85)
x_right = Inches(7.0)
y = Inches(2.2)

for i, (num, title) in enumerate(contents):
    col = i // 7
    row = i % 7
    sx = x_left if col == 0 else x_right
    sy = y + Inches(0.6) * row
    add_round(s, sx, sy, Inches(5.5), Inches(0.5), LIGHT_BG, line_color=BORDER)
    add_text(s, sx + Inches(0.2), sy + Inches(0.07), Inches(0.7), Inches(0.4),
             num, font_size=14, bold=True, color=ACCENT,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, sx + Inches(0.9), sy + Inches(0.07), Inches(4.4), Inches(0.4),
             title, font_size=13, color=TEXT_DARK, anchor=MSO_ANCHOR.MIDDLE)

add_footer(s, 2, TOTAL)

# ==================== СЛАЙД 3: АКТУАЛЬНОСТЬ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "01  ВВЕДЕНИЕ", "Актуальность темы исследования")

# Левая колонка - текст
add_text(s, Inches(0.85), Inches(2.2), Inches(7), Inches(0.5),
         "Цифровая трансформация правового поля",
         font_size=18, bold=True, color=MGIMO_BLUE)

text_block = (
    "Современное российское бизнес-сообщество сталкивается с растущей "
    "сложностью правового регулирования. По данным Минэкономразвития, "
    "ежегодно принимается более 500 нормативных актов, влияющих на "
    "деятельность МСП.\n\n"
    "При этом 73% субъектов малого бизнеса не имеют возможности содержать "
    "штатного юриста, а услуги внешних консультантов являются недоступными "
    "для большинства предпринимателей.\n\n"
    "Развитие технологий искусственного интеллекта открывает принципиально "
    "новые возможности для автоматизации юридических услуг и обеспечения "
    "правовой поддержки бизнеса."
)

add_text(s, Inches(0.85), Inches(2.85), Inches(7), Inches(4),
         text_block, font_size=12, color=TEXT_DARK)

# Правая колонка - ключевые цифры
add_round(s, Inches(8.3), Inches(2.2), Inches(4.5), Inches(4.7), MGIMO_BLUE)
add_text(s, Inches(8.5), Inches(2.4), Inches(4.1), Inches(0.4),
         "Ключевые показатели",
         font_size=14, bold=True, color=ACCENT)

stats = [
    ("6,2 млн", "субъектов МСП в РФ"),
    ("73%", "без штатного юриста"),
    ("450 млрд ₽", "объём рынка юр. услуг"),
    ("28%", "годовой рост LegalTech"),
    ("500+", "новых НПА в год"),
]

for i, (val, lbl) in enumerate(stats):
    ty = Inches(2.95) + Inches(0.75) * i
    add_text(s, Inches(8.5), ty, Inches(2.0), Inches(0.4),
             val, font_size=18, bold=True, color=ACCENT)
    add_text(s, Inches(8.5), ty + Inches(0.4), Inches(4.1), Inches(0.3),
             lbl, font_size=10, color=WHITE)

add_footer(s, 3, TOTAL)

# ==================== СЛАЙД 4: ЦЕЛЬ И ЗАДАЧИ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "02  ЦЕЛЬ И ЗАДАЧИ", "Что должно быть выполнено в рамках работы")

# Цель
add_round(s, Inches(0.6), Inches(2.1), Inches(12.1), Inches(1.3), MGIMO_BLUE)
add_text(s, Inches(0.85), Inches(2.25), Inches(2), Inches(0.4),
         "ЦЕЛЬ:", font_size=14, bold=True, color=ACCENT)
add_text(s, Inches(0.85), Inches(2.65), Inches(11.5), Inches(0.7),
         "Разработать функционирующий прототип цифровой платформы правовой "
         "поддержки на основе AI,",
         font_size=15, color=WHITE)
add_text(s, Inches(0.85), Inches(3.0), Inches(11.5), Inches(0.4),
         "обеспечивающей доступную юридическую помощь для субъектов МСП.",
         font_size=15, color=WHITE)

# Задачи - 6 в 2 ряда по 3
add_text(s, Inches(0.85), Inches(3.7), Inches(11), Inches(0.4),
         "ЗАДАЧИ:", font_size=14, bold=True, color=MGIMO_BLUE)

tasks = [
    ("1", "Проанализировать рынок\nюр. услуг для МСП"),
    ("2", "Спроектировать\nархитектуру системы"),
    ("3", "Реализовать\nAI-анализ документов"),
    ("4", "Интегрировать\nвнешние API (DaData)"),
    ("5", "Разработать алгоритм\nскоринга контрагентов"),
    ("6", "Развернуть систему\nв облачной инфраструктуре"),
]

x = Inches(0.6)
y = Inches(4.3)
tw = Inches(3.95)
th = Inches(1.3)
gap = Inches(0.1)

for i, (num, txt) in enumerate(tasks):
    row = i // 3
    col = i % 3
    sx = x + (tw + gap) * col
    sy = y + (th + Inches(0.15)) * row
    add_round(s, sx, sy, tw, th, LIGHT_BG, line_color=BORDER)
    add_oval(s, sx + Inches(0.2), sy + Inches(0.35), Inches(0.6),
             Inches(0.6), MGIMO_BLUE)
    add_text(s, sx + Inches(0.2), sy + Inches(0.37), Inches(0.6), Inches(0.6),
             num, font_size=20, bold=True, color=ACCENT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, sx + Inches(1.0), sy + Inches(0.25), tw - Inches(1.1),
             Inches(0.9), txt, font_size=12, color=TEXT_DARK,
             anchor=MSO_ANCHOR.MIDDLE)

add_footer(s, 4, TOTAL)

# ==================== СЛАЙД 5: ОБЪЕКТ, ПРЕДМЕТ, МЕТОДОЛОГИЯ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "03  МЕТОДОЛОГИЯ", "Объект и предмет исследования, методы")

# Три блока в столбик
blocks = [
    ("Объект исследования",
     "Процессы цифровизации юридических услуг для субъектов малого и среднего "
     "предпринимательства в Российской Федерации.",
     MGIMO_BLUE),
    ("Предмет исследования",
     "Программно-технические средства автоматизации правовой поддержки "
     "бизнеса с применением технологий искусственного интеллекта и обработки "
     "естественного языка.",
     MGIMO_LIGHT),
    ("Методы исследования",
     "Системный анализ, сравнительный анализ существующих решений, "
     "проектирование информационных систем, прототипирование, "
     "программная инженерия, экспериментальное тестирование.",
     ACCENT),
]

x = Inches(0.6)
y = Inches(2.2)
bw = Inches(12.1)
bh = Inches(1.4)
gap = Inches(0.2)

for i, (title, desc, color) in enumerate(blocks):
    sy = y + (bh + gap) * i
    add_round(s, x, sy, bw, bh, LIGHT_BG, line_color=BORDER)
    # Цветная полоска слева
    add_rect(s, x, sy, Inches(0.15), bh, color)
    add_text(s, x + Inches(0.4), sy + Inches(0.2), Inches(11.5), Inches(0.4),
             title, font_size=15, bold=True, color=color)
    add_text(s, x + Inches(0.4), sy + Inches(0.65), Inches(11.5), Inches(0.7),
             desc, font_size=12, color=TEXT_DARK)

add_footer(s, 5, TOTAL)

# ==================== СЛАЙД 6: АНАЛИЗ ПРЕДМЕТНОЙ ОБЛАСТИ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "04  АНАЛИЗ ПРЕДМЕТНОЙ ОБЛАСТИ", "Существующие решения и их недостатки")

# Таблица сравнения
columns = ["Платформа", "Функционал", "Целевая аудитория", "Недостатки"]
rows = [
    ("Контур.Фокус",
     "Проверка контрагентов",
     "Корпоративный сегмент",
     "Высокая цена, нет AI"),
    ("СПАРК",
     "Бизнес-разведка",
     "Крупные компании",
     "Не для МСП, дорого"),
    ("КонсультантПлюс",
     "База законов",
     "Юристы, бухгалтеры",
     "Нет автоматизации задач"),
    ("Правовед.ru",
     "Онлайн-консультации",
     "Физические лица",
     "Зависит от живого юриста"),
]

# Шапка
y = Inches(2.2)
col_widths = [Inches(2.5), Inches(3.5), Inches(3), Inches(3.1)]
x_positions = [Inches(0.6)]
for w in col_widths[:-1]:
    x_positions.append(x_positions[-1] + w)

for i, col in enumerate(columns):
    add_rect(s, x_positions[i], y, col_widths[i], Inches(0.55), MGIMO_BLUE,
             line_color=MGIMO_BLUE)
    add_text(s, x_positions[i], y + Inches(0.05), col_widths[i], Inches(0.5),
             col, font_size=12, bold=True, color=ACCENT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

for ri, vals in enumerate(rows):
    ry = y + Inches(0.55) + Inches(0.6) * ri
    bg = LIGHT_BG if ri % 2 == 0 else WHITE
    for ci, v in enumerate(vals):
        add_rect(s, x_positions[ci], ry, col_widths[ci], Inches(0.6), bg,
                 line_color=BORDER)
        add_text(s, x_positions[ci] + Inches(0.15), ry + Inches(0.05),
                 col_widths[ci] - Inches(0.3), Inches(0.5),
                 v, font_size=11, color=TEXT_DARK, anchor=MSO_ANCHOR.MIDDLE)

# Вывод
add_round(s, Inches(0.6), Inches(5.7), Inches(12.1), Inches(1.1),
          RGBColor(0xFF, 0xF8, 0xE7))
add_text(s, Inches(0.85), Inches(5.85), Inches(11.5), Inches(0.4),
         "Вывод по результатам анализа:", font_size=13, bold=True,
         color=MGIMO_BLUE)
add_text(s, Inches(0.85), Inches(6.25), Inches(11.5), Inches(0.5),
         "На рынке отсутствует комплексное решение, объединяющее AI-анализ "
         "документов, проверку контрагентов и генерацию документов в одном "
         "доступном продукте для МСП.",
         font_size=12, color=TEXT_DARK)

add_footer(s, 6, TOTAL)

# ==================== СЛАЙД 7: АРХИТЕКТУРА ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "05  АРХИТЕКТУРА СИСТЕМЫ", "Трёхзвенная архитектура клиент-сервер")

# Слой 1 - Клиент
add_round(s, Inches(0.6), Inches(2.2), Inches(12.1), Inches(0.85), MGIMO_BLUE)
add_text(s, Inches(0.8), Inches(2.3), Inches(2.5), Inches(0.4),
         "СЛОЙ 1: ПРЕДСТАВЛЕНИЕ", font_size=11, bold=True, color=ACCENT)
add_text(s, Inches(0.8), Inches(2.65), Inches(11.5), Inches(0.4),
         "Web-интерфейс (Next.js + React)  ·  Адаптивная вёрстка  ·  Авторизация пользователей",
         font_size=12, color=WHITE)

# Стрелки
arrow1 = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), Inches(3.15),
                            Inches(0.4), Inches(0.3))
arrow1.fill.solid()
arrow1.fill.fore_color.rgb = ACCENT
arrow1.line.fill.background()

# Слой 2 - Бизнес-логика
add_round(s, Inches(0.6), Inches(3.6), Inches(12.1), Inches(1.7), MGIMO_LIGHT)
add_text(s, Inches(0.8), Inches(3.7), Inches(3), Inches(0.4),
         "СЛОЙ 2: БИЗНЕС-ЛОГИКА", font_size=11, bold=True, color=ACCENT)

# 4 модуля
modules = [
    ("AI-анализ", "Claude/Gemini\nАнализ документов"),
    ("Контрагенты", "Скоринг\nDaData/ЕГРЮЛ"),
    ("Генерация", "AI-шаблоны\nдоговоров"),
    ("База права", "ГК, ТК, НК\nПоиск"),
]
mw = Inches(2.85)
mx = Inches(0.8)
my = Inches(4.15)
gap = Inches(0.1)
for i, (name, desc) in enumerate(modules):
    msx = mx + (mw + gap) * i
    add_round(s, msx, my, mw, Inches(1.05), WHITE)
    add_text(s, msx, my + Inches(0.1), mw, Inches(0.4),
             name, font_size=13, bold=True, color=MGIMO_BLUE,
             align=PP_ALIGN.CENTER)
    add_text(s, msx, my + Inches(0.5), mw, Inches(0.5),
             desc, font_size=10, color=TEXT_GREY, align=PP_ALIGN.CENTER)

# Стрелка вниз
arrow2 = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), Inches(5.4),
                            Inches(0.4), Inches(0.3))
arrow2.fill.solid()
arrow2.fill.fore_color.rgb = ACCENT
arrow2.line.fill.background()

# Слой 3 - Данные
add_round(s, Inches(0.6), Inches(5.85), Inches(12.1), Inches(1.05), TEXT_DARK)
add_text(s, Inches(0.8), Inches(5.95), Inches(2.5), Inches(0.4),
         "СЛОЙ 3: ДАННЫЕ", font_size=11, bold=True, color=ACCENT)

dbs = ["PostgreSQL (Neon)", "Внешние API: DaData", "AI: Claude · Gemini · Groq",
       "Хранилище файлов"]
for i, db in enumerate(dbs):
    add_text(s, Inches(0.8) + Inches(2.95) * i, Inches(6.4), Inches(2.9),
             Inches(0.4), "▸ " + db, font_size=11, color=WHITE)

add_footer(s, 7, TOTAL)

# ==================== СЛАЙД 8: ТЕХ. СТЕК ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "06  ТЕХНОЛОГИЧЕСКИЙ СТЕК", "Выбор технологий и его обоснование")

stack = [
    ("Frontend",
     "Next.js 16 (React 19, TypeScript)",
     "SSR, оптимизация SEO, типобезопасность"),
    ("Backend",
     "Next.js API Routes",
     "Единый язык TypeScript, simplicity, производительность"),
    ("ORM / БД",
     "Prisma + PostgreSQL (Neon)",
     "Типобезопасность запросов, serverless PostgreSQL"),
    ("Аутентификация",
     "NextAuth.js (JWT)",
     "Безопасность, поддержка OAuth провайдеров"),
    ("AI-движок",
     "Claude 4.5, Gemini 2.0, Groq",
     "Резервирование, качественный анализ юр. текстов"),
    ("Внешние API",
     "DaData (ЕГРЮЛ, ОКВЭД, фин. данные)",
     "Официальный источник данных по компаниям РФ"),
    ("Деплой",
     "Vercel + GitHub Actions",
     "Continuous Deployment, CDN, edge functions"),
]

y = Inches(2.2)
for i, (cat, tech, why) in enumerate(stack):
    ty = y + Inches(0.65) * i
    add_round(s, Inches(0.6), ty, Inches(12.1), Inches(0.55), LIGHT_BG,
              line_color=BORDER)
    add_text(s, Inches(0.85), ty + Inches(0.1), Inches(2.0), Inches(0.4),
             cat, font_size=12, bold=True, color=MGIMO_BLUE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(2.95), ty + Inches(0.1), Inches(4.2), Inches(0.4),
             tech, font_size=11, bold=True, color=TEXT_DARK,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(7.3), ty + Inches(0.1), Inches(5.3), Inches(0.4),
             "→ " + why, font_size=11, color=TEXT_GREY,
             anchor=MSO_ANCHOR.MIDDLE, italic=True)

add_footer(s, 8, TOTAL)

# ==================== СЛАЙД 9: AI-АНАЛИЗ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "07  РЕАЛИЗАЦИЯ", "Модуль AI-анализа документов")

# Описание
add_text(s, Inches(0.85), Inches(2.2), Inches(7.5), Inches(0.5),
         "Принцип работы:", font_size=16, bold=True, color=MGIMO_BLUE)

steps = [
    "Пользователь загружает документ (PDF, DOCX)",
    "Система извлекает текст через pdf-parse / mammoth",
    "Текст передаётся в LLM с экспертным промптом",
    "AI выявляет: риски, неоднозначные формулировки, отсутствие пунктов",
    "Привязка найденных моментов к статьям ГК РФ из базы знаний",
    "Формирование структурированного отчёта с рекомендациями",
]

for i, step in enumerate(steps):
    ty = Inches(2.8) + Inches(0.5) * i
    add_oval(s, Inches(0.95), ty + Inches(0.1), Inches(0.25), Inches(0.25),
             ACCENT)
    add_text(s, Inches(0.95), ty + Inches(0.08), Inches(0.25), Inches(0.25),
             str(i + 1), font_size=10, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.35), ty + Inches(0.05), Inches(7), Inches(0.4),
             step, font_size=12, color=TEXT_DARK)

# Правая панель - технические детали
add_round(s, Inches(8.6), Inches(2.2), Inches(4.2), Inches(4.7), MGIMO_BLUE)
add_text(s, Inches(8.8), Inches(2.4), Inches(3.8), Inches(0.4),
         "Технические детали", font_size=14, bold=True, color=ACCENT)

details = [
    ("Размер промпта", "~3 000 токенов"),
    ("Модель LLM", "Claude 4.5 Sonnet"),
    ("Резерв", "Gemini 2.0 / Groq"),
    ("Среднее время", "~15 секунд"),
    ("Точность анализа", "Верифицируется БД"),
    ("Форматы файлов", "PDF, DOCX, TXT"),
    ("Макс. размер", "10 МБ / документ"),
    ("Контекст", "до 200К токенов"),
]

for i, (lbl, val) in enumerate(details):
    ty = Inches(3.0) + Inches(0.45) * i
    add_text(s, Inches(8.8), ty, Inches(2.0), Inches(0.4),
             lbl, font_size=11, color=WHITE)
    add_text(s, Inches(10.8), ty, Inches(1.9), Inches(0.4),
             val, font_size=11, bold=True, color=ACCENT, align=PP_ALIGN.RIGHT)

add_footer(s, 9, TOTAL)

# ==================== СЛАЙД 10: КОНТРАГЕНТЫ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "08  РЕАЛИЗАЦИЯ", "Модуль проверки контрагентов")

# Слева - источники данных
add_text(s, Inches(0.85), Inches(2.2), Inches(6), Inches(0.5),
         "Источники данных:", font_size=16, bold=True, color=MGIMO_BLUE)

sources = [
    ("DaData API", "Официальные данные ЕГРЮЛ, статус, ОКВЭД, капитал"),
    ("ЕГРЮЛ ФНС", "Резервный источник базовой информации"),
    ("КАД (картотека судов)", "Активные и завершённые судебные дела"),
    ("ФССП", "Информация о задолженностях, исполнительных производствах"),
]

for i, (src, desc) in enumerate(sources):
    ty = Inches(2.85) + Inches(0.85) * i
    add_round(s, Inches(0.85), ty, Inches(6), Inches(0.75), LIGHT_BG,
              line_color=BORDER)
    add_text(s, Inches(1.05), ty + Inches(0.1), Inches(5.7), Inches(0.35),
             src, font_size=12, bold=True, color=MGIMO_BLUE)
    add_text(s, Inches(1.05), ty + Inches(0.4), Inches(5.7), Inches(0.4),
             desc, font_size=10, color=TEXT_GREY)

# Справа - алгоритм скоринга
add_round(s, Inches(7.05), Inches(2.2), Inches(5.7), Inches(4.7), MGIMO_BLUE)
add_text(s, Inches(7.25), Inches(2.4), Inches(5.3), Inches(0.5),
         "Алгоритм скоринга рисков", font_size=14, bold=True, color=ACCENT)

factors = [
    ("Возраст компании", "+/- 15 баллов"),
    ("Текущий статус (ЕГРЮЛ)", "+/- 30 баллов"),
    ("Активные судебные дела", "- 5 за каждое"),
    ("Завершённые проигрыши", "- 3 за каждое"),
    ("Долги по ФССП", "- 20 баллов"),
    ("Размер уставного капитала", "+/- 10 баллов"),
    ("ОКВЭД (вид деятельности)", "± 5 баллов"),
    ("Финансовые показатели", "± 15 баллов"),
]

for i, (f, w) in enumerate(factors):
    ty = Inches(3.0) + Inches(0.4) * i
    add_text(s, Inches(7.4), ty, Inches(3), Inches(0.3),
             "•  " + f, font_size=11, color=WHITE)
    add_text(s, Inches(10.6), ty, Inches(2), Inches(0.3),
             w, font_size=10, color=ACCENT, align=PP_ALIGN.RIGHT)

# Уровни риска
add_text(s, Inches(7.25), Inches(6.3), Inches(5.3), Inches(0.4),
         "Итоговая оценка: НИЗКИЙ / СРЕДНИЙ / ВЫСОКИЙ риск",
         font_size=11, italic=True, color=ACCENT, align=PP_ALIGN.CENTER)

add_footer(s, 10, TOTAL)

# ==================== СЛАЙД 11: БАЗА ПРАВА ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "09  РЕАЛИЗАЦИЯ", "База права и генерация документов")

# Левая часть - база права
add_round(s, Inches(0.6), Inches(2.2), Inches(6.05), Inches(4.65), LIGHT_BG,
          line_color=BORDER)
add_rect(s, Inches(0.6), Inches(2.2), Inches(6.05), Inches(0.5), MGIMO_BLUE)
add_text(s, Inches(0.6), Inches(2.25), Inches(6.05), Inches(0.4),
         "База российского права", font_size=14, bold=True, color=ACCENT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

codes = [
    ("ГК РФ", "Гражданский кодекс"),
    ("ТК РФ", "Трудовой кодекс"),
    ("НК РФ", "Налоговый кодекс"),
    ("КоАП", "Об административных нарушениях"),
    ("ФЗ", "Федеральные законы"),
    ("Постановления", "Правительства и Пленумов"),
]

for i, (code, name) in enumerate(codes):
    ty = Inches(2.95) + Inches(0.6) * i
    add_round(s, Inches(0.85), ty, Inches(5.55), Inches(0.5), WHITE,
              line_color=BORDER)
    add_text(s, Inches(1.05), ty + Inches(0.07), Inches(1.5), Inches(0.4),
             code, font_size=13, bold=True, color=MGIMO_BLUE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(2.55), ty + Inches(0.07), Inches(3.8), Inches(0.4),
             name, font_size=11, color=TEXT_DARK, anchor=MSO_ANCHOR.MIDDLE)

# Правая часть - генерация
add_round(s, Inches(6.85), Inches(2.2), Inches(6.05), Inches(4.65), LIGHT_BG,
          line_color=BORDER)
add_rect(s, Inches(6.85), Inches(2.2), Inches(6.05), Inches(0.5), MGIMO_BLUE)
add_text(s, Inches(6.85), Inches(2.25), Inches(6.05), Inches(0.4),
         "Генерация документов", font_size=14, bold=True, color=ACCENT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

docs = [
    ("Договор оказания услуг", "Стандартный + кастомизация"),
    ("Договор поставки", "С учётом отрасли"),
    ("Трудовой договор", "По ТК РФ"),
    ("Исковое заявление", "Для арбитража"),
    ("Претензионное письмо", "Досудебное урегулирование"),
    ("Доверенность", "Простая письменная"),
]

for i, (doc, desc) in enumerate(docs):
    ty = Inches(2.95) + Inches(0.6) * i
    add_round(s, Inches(7.1), ty, Inches(5.55), Inches(0.5), WHITE,
              line_color=BORDER)
    add_text(s, Inches(7.3), ty + Inches(0.05), Inches(3.5), Inches(0.4),
             doc, font_size=11, bold=True, color=MGIMO_BLUE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(10.0), ty + Inches(0.05), Inches(2.5), Inches(0.4),
             desc, font_size=9, color=TEXT_GREY, italic=True,
             anchor=MSO_ANCHOR.MIDDLE)

add_footer(s, 11, TOTAL)

# ==================== СЛАЙД 12: ДЕПЛОЙ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "10  ДЕПЛОЙ И ИНФРАСТРУКТУРА", "Развёртывание системы в облаке")

# Схема деплоя
add_text(s, Inches(0.85), Inches(2.2), Inches(12), Inches(0.5),
         "Облачная архитектура продакшена:",
         font_size=16, bold=True, color=MGIMO_BLUE)

# Блоки инфраструктуры
infra = [
    ("GitHub", "Версионный\nконтроль кода",
     "→", LIGHT_BG, MGIMO_BLUE),
    ("Vercel", "Хостинг\nNext.js приложения",
     "→", MGIMO_BLUE, WHITE),
    ("Neon", "PostgreSQL\nбаза данных",
     "↔", LIGHT_BG, MGIMO_BLUE),
    ("DaData", "API\nданных компаний",
     "↔", LIGHT_BG, MGIMO_BLUE),
    ("LLM API", "AI-провайдеры\n(Claude/Gemini)",
     "", MGIMO_BLUE, WHITE),
]

x = Inches(0.85)
y = Inches(3.0)
iw = Inches(2.2)
ih = Inches(1.6)
gap = Inches(0.15)

for i, (name, desc, arr, bg, txt) in enumerate(infra):
    sx = x + (iw + gap) * i
    add_round(s, sx, y, iw, ih, bg, line_color=BORDER)
    add_text(s, sx, y + Inches(0.25), iw, Inches(0.5),
             name, font_size=18, bold=True,
             color=ACCENT if bg == MGIMO_BLUE else MGIMO_BLUE,
             align=PP_ALIGN.CENTER)
    add_text(s, sx + Inches(0.15), y + Inches(0.85), iw - Inches(0.3),
             Inches(0.7), desc, font_size=11, color=txt, align=PP_ALIGN.CENTER)

# Метрики деплоя
add_text(s, Inches(0.85), Inches(5.0), Inches(12), Inches(0.4),
         "Производственные характеристики:",
         font_size=14, bold=True, color=MGIMO_BLUE)

metrics = [
    ("99,9%", "SLA доступности"),
    ("< 200мс", "Время отклика API"),
    ("Auto", "Масштабирование"),
    ("HTTPS", "Шифрование"),
    ("CI/CD", "Деплой за 90 сек"),
]

x = Inches(0.85)
y = Inches(5.5)
mw = Inches(2.35)
gap = Inches(0.15)

for i, (val, lbl) in enumerate(metrics):
    sx = x + (mw + gap) * i
    add_round(s, sx, y, mw, Inches(1.2), MGIMO_BLUE)
    add_text(s, sx, y + Inches(0.15), mw, Inches(0.5),
             val, font_size=20, bold=True, color=ACCENT,
             align=PP_ALIGN.CENTER)
    add_text(s, sx, y + Inches(0.7), mw, Inches(0.4),
             lbl, font_size=11, color=WHITE, align=PP_ALIGN.CENTER)

add_footer(s, 12, TOTAL)

# ==================== СЛАЙД 13: РЕЗУЛЬТАТЫ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "11  РЕЗУЛЬТАТЫ И ДОСТИЖЕНИЯ", "Что удалось реализовать в рамках работы")

# Главное достижение
add_round(s, Inches(0.6), Inches(2.2), Inches(12.1), Inches(0.95), GREEN)
add_text(s, Inches(0.6), Inches(2.3), Inches(12.1), Inches(0.85),
         "✓  Разработан и развёрнут в продакшене работающий прототип "
         "AI-платформы правовой поддержки",
         font_size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)

# 6 модулей
results = [
    ("AI-анализ документов",
     "Реализован модуль автоматического анализа договоров с выявлением "
     "юридических рисков и привязкой к статьям ГК РФ"),
    ("Проверка контрагентов",
     "Интегрирован API DaData, разработан алгоритм скоринга на основе "
     "8 факторов с 3 уровнями риска"),
    ("База знаний по праву",
     "Создана структурированная БД ключевых статей кодексов РФ с "
     "полнотекстовым поиском"),
    ("Генерация документов",
     "Реализована система AI-генерации договоров, исковых заявлений и "
     "претензий по запросу"),
    ("Аутентификация и тарифы",
     "JWT-авторизация, поддержка Google OAuth, разделение функционала "
     "по 3 тарифным планам"),
    ("Облачный деплой",
     "Приложение развёрнуто на Vercel с PostgreSQL базой Neon, настроен "
     "CI/CD через GitHub"),
]

x = Inches(0.6)
y = Inches(3.4)
rw = Inches(3.95)
rh = Inches(1.65)
gap = Inches(0.1)

for i, (title, desc) in enumerate(results):
    row = i // 3
    col = i % 3
    sx = x + (rw + gap) * col
    sy = y + (rh + gap) * row
    add_round(s, sx, sy, rw, rh, LIGHT_BG, line_color=BORDER)
    add_oval(s, sx + Inches(0.2), sy + Inches(0.25), Inches(0.4),
             Inches(0.4), GREEN)
    add_text(s, sx + Inches(0.2), sy + Inches(0.27), Inches(0.4), Inches(0.4),
             "✓", font_size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, sx + Inches(0.7), sy + Inches(0.2), rw - Inches(0.85),
             Inches(0.45), title, font_size=13, bold=True, color=MGIMO_BLUE)
    add_text(s, sx + Inches(0.25), sy + Inches(0.65), rw - Inches(0.4),
             Inches(0.95), desc, font_size=10, color=TEXT_GREY)

add_footer(s, 13, TOTAL)

# ==================== СЛАЙД 14: СЛОЖНОСТИ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "12  СЛОЖНОСТИ РАЗРАБОТКИ", "С какими проблемами пришлось столкнуться")

challenges = [
    ("Точность AI-анализа",
     "Модели LLM могут ошибаться в специфических юридических вопросах",
     "Расширенные промпты, верификация по БД, дисклеймеры для пользователя"),
    ("Различия SQLite и PostgreSQL",
     "Локальная разработка велась на SQLite, а продакшн использует PostgreSQL",
     "Унификация запросов через Prisma ORM, создание миграций для PG"),
    ("Интеграция с DaData",
     "Кэширование старых данных приводило к показу демо-данных вместо реальных",
     "Введение forceRefresh параметра, проверка флага dataSource в БД"),
    ("Производительность LLM",
     "Запросы к Claude занимали 10-15 секунд, что замедляло UX",
     "Поддержка нескольких провайдеров (Groq для скорости), стриминг ответов"),
    ("Ограничения бесплатных API",
     "DaData имеет лимит запросов на бесплатном тарифе",
     "Кэширование результатов, обновление профилей раз в 30 дней"),
]

y = Inches(2.1)
for i, (name, problem, solution) in enumerate(challenges):
    sy = y + Inches(0.93) * i
    add_round(s, Inches(0.6), sy, Inches(12.1), Inches(0.85), LIGHT_BG,
              line_color=BORDER)
    add_rect(s, Inches(0.6), sy, Inches(0.12), Inches(0.85), ACCENT)
    add_text(s, Inches(0.85), sy + Inches(0.08), Inches(11.5), Inches(0.35),
             name, font_size=12, bold=True, color=MGIMO_BLUE)
    add_text(s, Inches(0.85), sy + Inches(0.4), Inches(5.5), Inches(0.4),
             "Проблема: " + problem, font_size=10, color=TEXT_GREY)
    add_text(s, Inches(6.5), sy + Inches(0.4), Inches(6), Inches(0.4),
             "→ Решение: " + solution, font_size=10, color=GREEN)

add_footer(s, 14, TOTAL)

# ==================== СЛАЙД 15: ПЕРСПЕКТИВЫ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "13  ПЕРСПЕКТИВЫ РАЗВИТИЯ", "Дальнейшие направления работы над проектом")

perspectives = [
    ("Краткосрочные (3-6 мес)",
     ["Расширение базы права до 50+ статей основных кодексов",
      "Telegram-бот для быстрых юридических консультаций",
      "Поддержка распознавания сканированных документов (Vision AI)",
      "Интеграция с электронной подписью"], MGIMO_BLUE),
    ("Среднесрочные (6-12 мес)",
     ["Mobile-приложение для iOS и Android",
      "API для интеграции в CRM/ERP системы (1С, Битрикс24)",
      "Система коллаборации (комментарии, версии документов)",
      "Расширенная аналитика и отчёты"], MGIMO_LIGHT),
    ("Долгосрочные (1-3 года)",
     ["Выход на рынки СНГ (Казахстан, Беларусь)",
      "Голосовой AI-ассистент для устных консультаций",
      "Прогнозирование исхода судебных дел",
      "Корпоративные лицензии для юридических департаментов"], ACCENT),
]

x = Inches(0.6)
y = Inches(2.2)
pw = Inches(3.95)
ph = Inches(4.6)
gap = Inches(0.1)

for i, (period, items, color) in enumerate(perspectives):
    sx = x + (pw + gap) * i
    add_round(s, sx, y, pw, ph, LIGHT_BG, line_color=BORDER)
    add_rect(s, sx, y, pw, Inches(0.7), color)
    add_text(s, sx, y + Inches(0.15), pw, Inches(0.5),
             period, font_size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for j, item in enumerate(items):
        ty = y + Inches(1.0) + Inches(0.85) * j
        add_text(s, sx + Inches(0.2), ty, Inches(0.2), Inches(0.4),
                 "▸", font_size=14, bold=True, color=color)
        add_text(s, sx + Inches(0.45), ty - Inches(0.05), pw - Inches(0.6),
                 Inches(0.85), item, font_size=11, color=TEXT_DARK)

add_footer(s, 15, TOTAL)

# ==================== СЛАЙД 16: ЗАКЛЮЧЕНИЕ ====================
s = prs.slides.add_slide(BLANK)
add_bg(s)
header(s, "14  ЗАКЛЮЧЕНИЕ", "Основные выводы по результатам работы")

# Основной текст
conclusions = [
    "В ходе курсовой работы был проведён анализ рынка правовых услуг "
    "для МСП в России, выявлены ключевые проблемы и определена ниша "
    "для AI-решения.",
    "Спроектирована и реализована трёхзвенная архитектура цифровой "
    "платформы с применением современного стека: Next.js 16, Prisma, "
    "PostgreSQL, NextAuth.js.",
    "Интегрированы три ведущих AI-провайдера (Claude, Gemini, Groq) и "
    "официальный API сервиса DaData для получения данных о российских "
    "юридических лицах.",
    "Разработан оригинальный алгоритм скоринга контрагентов на основе "
    "8 факторов с тремя уровнями риска.",
    "Прототип успешно развёрнут в продакшен-окружении на платформе "
    "Vercel с облачной БД Neon и доступен по публичному URL.",
]

for i, text in enumerate(conclusions):
    ty = Inches(2.2) + Inches(0.7) * i
    add_oval(s, Inches(0.85), ty + Inches(0.15), Inches(0.3), Inches(0.3),
             ACCENT)
    add_text(s, Inches(0.85), ty + Inches(0.13), Inches(0.3), Inches(0.3),
             str(i + 1), font_size=12, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.3), ty + Inches(0.05), Inches(11.4), Inches(0.7),
             text, font_size=12, color=TEXT_DARK)

# Финальная плашка
add_round(s, Inches(0.6), Inches(6.0), Inches(12.1), Inches(0.85), MGIMO_BLUE)
add_text(s, Inches(0.6), Inches(6.05), Inches(12.1), Inches(0.4),
         "Цели и задачи курсовой работы выполнены в полном объёме.",
         font_size=14, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(s, Inches(0.6), Inches(6.45), Inches(12.1), Inches(0.4),
         "Спасибо за внимание!",
         font_size=14, color=WHITE, align=PP_ALIGN.CENTER, italic=True)

add_footer(s, 16, TOTAL)

prs.save("/home/user/don/presentations/02_Курсовой_отчет_Хадызов.pptx")
print("Курсовая презентация создана: 16 слайдов")
